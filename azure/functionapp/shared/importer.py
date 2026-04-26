"""
Parse Sheet1 of the SIT Skate Club Rules Excel file and upsert into DynamoDB
members table. Ported near-unchanged from Phase 3 — openpyxl works in Lambda.

Sheet1 column layout (0-indexed, confirmed from actual file):
  0  Id
  1  Start time
  2  Completion time
  3  Email
  4  Name                  -- always None (form artifact), skip
  5  Full Name:            -- actual member name
  6  Telegram Handle:
  7  SIT Student/alumni?   ("Yes" / "No")
  8  SIT Student ID        (blank for non-SIT members)
  9  Course: (short)       -- skipped, we use Full Course Name
  10 Full Course Name      (blank for non-SIT members)
  11 Cluster               (blank for non-SIT members)
  12 Year                  (blank for non-SIT members)
  13 NRIC masked           -- not stored
  14 Masked NRIC           -- not stored
  15 Acknowledgement text  -- not stored
"""
from __future__ import annotations

from datetime import datetime, timezone

import openpyxl

from shared import db


def _cell(row, idx):
    val = row[idx].value
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def import_from_file(path: str) -> dict:
    """Read the xlsx and upsert members. Returns summary counts."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]

    now = datetime.now(timezone.utc).isoformat()
    counts = {"new": 0, "updated": 0, "unchanged": 0, "skipped": 0}
    rows_seen = 0

    for i, row in enumerate(ws.iter_rows()):
        if i == 0:
            continue
        rows_seen += 1

        full_name = _cell(row, 5)
        handle_raw = _cell(row, 6)
        if not full_name or not handle_raw:
            counts["skipped"] += 1
            continue

        username = handle_raw.lstrip("@").lower()
        is_sit = (_cell(row, 7) or "").lower() == "yes"
        student_id = _cell(row, 8) if is_sit else None
        full_course = _cell(row, 10)
        cluster = _cell(row, 11)
        year = _cell(row, 12)

        result = db.upsert_member(
            username=username,
            full_name=full_name,
            is_sit_student=is_sit,
            student_id=student_id,
            full_course_name=full_course,
            cluster=cluster,
            year=year,
            imported_at=now,
        )
        counts[result] += 1

    wb.close()
    counts["total"] = rows_seen
    return counts

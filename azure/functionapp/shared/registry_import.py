"""Header-based rules registry import. No NRIC or contact-email fields are stored."""
from __future__ import annotations

import re
from datetime import UTC, datetime

import openpyxl
from azure.data.tables import UpdateMode
from shared import db

ACKNOWLEDGEMENT = 'i have read and acknowledged the updated rules in sit inline skating club and will abide by them.'


def normalized(value):
    return ' '.join(str(value or '').strip().lower().split())


def parse_workbook(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['Sheet1'] if 'Sheet1' in workbook.sheetnames else workbook.worksheets[0]
        rows = iter(sheet.values)
        header_row = next(rows, None)
        if not header_row:
            raise ValueError('The registry workbook is empty.')
        headers = [normalized(v).rstrip(':') for v in header_row]
        def column(name, required=False):
            try:
                return headers.index(name)
            except ValueError:
                if required:
                    raise ValueError(f'Not a rules registry: missing {name} column.') from None
                return None
        name_col = column('full name', True)
        handle_col = column('telegram handle', True)
        ack_cols = [i for i, h in enumerate(headers) if 'acknowledg' in h or h == 'column']
        if len(ack_cols) != 1:
            raise ValueError('Expected one rules acknowledgement column; import was not applied.')
        result = []
        for number, row in enumerate(rows, start=2):
            def value(index):
                return str(row[index]).strip() if index is not None and index < len(row) and row[index] is not None else ''
            handle = value(handle_col).lstrip('@').lower()
            result.append({
                'row': number, 'username': handle, 'full_name': value(name_col),
                'valid': bool(re.fullmatch(r'[a-zA-Z][a-zA-Z0-9_]{4,31}', handle) and value(name_col)),
                'rules_acknowledged': normalized(value(ack_cols[0])) == ACKNOWLEDGEMENT,
                'source_response_id': value(column('id')),
                'is_sit_student': normalized(value(column('sit student/alumni?'))) == 'yes',
                'student_id': value(column('sit student id')),
                'full_course_name': value(column('full course name')),
                'cluster': value(column('cluster')), 'year': value(column('year')),
            })
        identities = {}
        for item in result:
            identities.setdefault(item['username'], set()).add((normalized(item['full_name']), item['student_id']))
        for item in result:
            item['ambiguous'] = len(identities[item['username']]) > 1
        return result
    finally:
        workbook.close()


def import_from_file(path):
    rows = parse_workbook(path)
    now = datetime.now(UTC).isoformat()
    counts = dict(new=0, updated=0, unchanged=0, skipped=0, total=len(rows), acknowledged=0, ambiguous=0)
    for row in rows:
        if not row['valid']:
            counts['skipped'] += 1
            continue
        fields = {k: row[k] for k in (
            'username', 'full_name', 'is_sit_student', 'student_id', 'full_course_name', 'cluster', 'year',
            'rules_acknowledged', 'ambiguous', 'source_response_id')}
        existing = db.get_member(row['username'])
        identity_conflict = bool(existing and (existing.get('identity_conflict') or
            normalized(existing.get('full_name')) != normalized(row['full_name']) or
            normalized(existing.get('student_id')) != normalized(row['student_id'])))
        fields['identity_conflict'] = identity_conflict
        fields['ambiguous'] = row['ambiguous'] or identity_conflict
        outcome = 'new' if existing is None else ('updated' if any(existing.get(k) != v for k, v in fields.items()) else 'unchanged')
        # One write prevents new identity data retaining stale acknowledgement
        # metadata if a request fails partway through importing this record.
        db._members.upsert_entity({
            'PartitionKey': 'MEMBER', 'RowKey': row['username'],
            **fields, 'imported_at': now, 'rules_imported_at': now,
        }, mode=UpdateMode.MERGE)
        counts[outcome] += 1
        counts['acknowledged'] += int(fields['rules_acknowledged'] and not fields['ambiguous'])
        counts['ambiguous'] += int(fields['ambiguous'])
    return counts
